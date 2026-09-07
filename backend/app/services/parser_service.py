import base64
from io import BytesIO
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

class ParserService:
    @staticmethod
    def get_extension(filename):
        return filename.rsplit(".", 1)[1].lower() if "." in filename else ""

    @staticmethod
    def extract_document_text(file, extension):
        """Extract readable text from PDF, DOCX, TXT, or PPTX."""
        if extension == "txt":
            return file.read().decode("utf-8", errors="ignore")

        if extension == "pdf":
            reader = PdfReader(file)
            pages = []
            for idx, page in enumerate(reader.pages, start=1):
                text = page.extract_text()
                if text and text.strip():
                    pages.append(f"--- Page {idx} ---\n" + text.strip())
            return "\n\n".join(pages)

        if extension == "docx":
            doc = Document(file)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)

        if extension == "pptx":
            prs = Presentation(file)
            slides = []
            for idx, slide in enumerate(prs.slides, start=1):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    slides.append(f"--- Slide {idx} ---\n" + "\n".join(slide_texts))
            return "\n\n".join(slides)

        raise ValueError(f"Unsupported document extension: {extension}")

    @staticmethod
    def encode_image(image_bytes):
        return base64.b64encode(image_bytes).decode("utf-8")
